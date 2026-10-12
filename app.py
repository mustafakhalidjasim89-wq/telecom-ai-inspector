import base64
import hashlib
import io
import json
import os
import re
import time
from datetime import datetime
from pathlib import Path

import fitz
import pandas as pd
import requests
import streamlit as st
from PIL import Image, ImageOps, ImageStat

LOCAL_MODEL = 'qwen2.5vl:7b'
CLOUD_MODEL = 'Qwen/Qwen2.5-VL-7B-Instruct'
CLOUD_DEPLOYMENT = True
CACHE = Path(__file__).parent / 'cache'
if not CLOUD_DEPLOYMENT:
    CACHE.mkdir(exist_ok=True)
CATEGORIES = ['Housekeeping', 'Power', 'RAN', 'Transmission', 'Tower', 'Civil', 'Safety', 'Other']
PROMPT = """You are a conservative telecom field photo inspection assistant. Return ONLY one valid JSON object, no markdown:
{"photo_distance":"close|medium|far|unknown","distance_reason":"","summary":"","defects":[{"category":"Power","defect":"","evidence":"","action":"","confidence":"high|medium|low","impact":"action_required|no_significant_impact|needs_verification","priority":"critical|high|medium|low","visibility":"clear|limited"}],"uncertain_observations":[]}
RULES:
1. Analyze only directly visible evidence in THIS image. Never infer loose terminals, live exposed wires, leakage, corrosion, missing equipment, bad electrical connections, dangerous faults or structural damage without clear visual proof.
2. Identify camera distance: FAR if telecom components are small or lack sufficient detail to verify their condition, even if the image resolution is high. For FAR photos, report only clearly visible LARGE-SCALE issues (e.g., vegetation obstructing access, fallen tower parts, obvious obstruction, standing water). Never diagnose small bolts, connections, cable terminations, equipment internals or minor rust from afar. Set visibility=limited and impact=needs_verification when uncertain.
3. Focus on items requiring an actual corrective action because of plausible operational, access, safety or equipment risk. For minor purely cosmetic items with no apparent impact, use impact=no_significant_impact, priority=low, action="No immediate action; routine monitoring" and keep the wording SHORT.
4. For a clear visible defect needing corrective action, use impact=action_required. If evidence is inadequate, use impact=needs_verification, action="Obtain close-up photo / site verification"; do NOT state a fault as proven.
5. Category must be exactly one of Housekeeping, Power, RAN, Transmission, Tower, Civil, Safety, Other. Categorize by affected equipment, not always Housekeeping.
6. Do not produce repetitive observations or generic advice. If no verifiable defect, defects=[].
7. Evidence must describe what is actually visible, not a conclusion. English, concise. All findings need engineer verification.
"""
NEGATIVE = re.compile(r'^(?:no |none|nothing |not |without |absence of |there (?:is|are) no |no visible |good condition|appears (?:fine|normal)|properly |adequate )', re.I)
SITE_RE = re.compile(r'\b(?:BAG|ANB|KUT|DIY|BQA|BSR|NJF|KRB|MOS)\d{3,5}\b', re.I)


def prepared(raw, max_side):
    with Image.open(io.BytesIO(raw)) as im:
        im = ImageOps.exif_transpose(im).convert('RGB')
        im.thumbnail((max_side, max_side))
        out = io.BytesIO()
        im.save(out, 'JPEG', quality=82)
        return out.getvalue()


def photo_quality(raw, min_side, min_kb):
    try:
        with Image.open(io.BytesIO(raw)) as im:
            if min(im.size) < min_side or len(raw) < min_kb * 1024:
                return False
            im.thumbnail((160, 160))
            return ImageStat.Stat(im.convert('L')).stddev[0] > 8
    except Exception:
        return False


def pdf_jobs(data, name, max_images, min_side, min_kb, render_pages, max_side):
    jobs, texts, notes, seen = [], [], [], set()
    with fitz.open(stream=data, filetype='pdf') as doc:
        for pi, page in enumerate(doc):
            txt = page.get_text('text')
            texts.append({'Source': name, 'Page': pi + 1, 'Text': txt[:30000]})
            if render_pages:
                raw_images = [(f'{name} | page {pi+1} (rendered)', page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False).tobytes('png'))]
            else:
                raw_images = []
                for ii, item in enumerate(page.get_images(full=True), 1):
                    try:
                        raw = doc.extract_image(item[0])['image']
                        if photo_quality(raw, min_side, min_kb):
                            raw_images.append((f'{name} | page {pi+1} image {ii}', raw))
                    except Exception:
                        continue
            for label, raw in raw_images:
                try:
                    key = hashlib.sha256(raw).hexdigest()
                    if key not in seen:
                        seen.add(key)
                        jobs.append((label, prepared(raw, max_side), pi + 1))
                except Exception:
                    continue
                if len(jobs) >= max_images:
                    break
            if len(jobs) >= max_images:
                notes.append(f'{name}: reached image limit {max_images}.')
                break
        notes.append(f'{name}: {len(doc)} pages; {len(jobs)} images queued; {len(texts)} pages of text extracted.')
    return jobs, texts, notes


def call_ollama(img, model, url, timeout):
    payload = {
        'model': model, 'stream': False, 'format': 'json',
        'options': {'temperature': 0, 'num_predict': 850},
        'messages': [
            {'role': 'system', 'content': PROMPT},
            {'role': 'user', 'content': 'Identify ONLY confirmed visible telecom site defects. JSON only.',
             'images': [base64.b64encode(img).decode('ascii')]},
        ],
    }
    response = requests.post(url.rstrip('/') + '/api/chat', json=payload, timeout=timeout)
    response.raise_for_status()
    data = response.json()
    return data.get('message', {}).get('content', ''), data.get('total_duration', 0) / 1e9


def call_runpod(img, model, base_url, api_key, timeout):
    # OpenAI-compatible vLLM API, without an additional Python SDK dependency.
    image_url = 'data:image/jpeg;base64,' + base64.b64encode(img).decode('ascii')
    payload = {
        'model': model,
        'temperature': 0,
        'max_tokens': 1100,
        'messages': [
            {'role': 'system', 'content': PROMPT},
            {'role': 'user', 'content': [
                {'type': 'text', 'text': 'Identify ONLY confirmed visible telecom site defects. Return valid JSON only.'},
                {'type': 'image_url', 'image_url': {'url': image_url}},
            ]},
        ],
    }
    headers = {'Content-Type': 'application/json', 'Authorization': f'Bearer {api_key}'}
    started = time.perf_counter()
    response = requests.post(base_url.rstrip('/') + '/chat/completions', headers=headers,
                             json=payload, timeout=timeout)
    response.raise_for_status()
    data = response.json()
    raw = data['choices'][0]['message'].get('content', '')
    if isinstance(raw, list):
        raw = ''.join(part.get('text', '') for part in raw if isinstance(part, dict))
    return str(raw), time.perf_counter() - started


def parse(raw):
    try:
        cleaned = re.sub(r'^```(?:json)?\s*|\s*```$', '', raw.strip(), flags=re.I)
        data = json.loads(cleaned)
        if not isinstance(data, dict) or not isinstance(data.get('defects'), list):
            return [], {}, 'Missing defects array'
        distance = str(data.get('photo_distance', 'unknown')).lower().strip()
        if distance not in ('close', 'medium', 'far', 'unknown'):
            distance = 'unknown'
        context = {'Photo Distance': distance,
                   'Distance Reason': str(data.get('distance_reason', ''))[:240],
                   'Uncertain Observations': '; '.join(map(str, data.get('uncertain_observations', [])))[:700]}
        out = []
        for d in data['defects']:
            if not isinstance(d, dict):
                continue
            defect = str(d.get('defect', '')).strip()
            evidence = str(d.get('evidence', '')).strip()
            if not defect or not evidence or NEGATIVE.match(defect) or NEGATIVE.match(evidence):
                continue
            category = str(d.get('category', 'Other')).strip()
            if category not in CATEGORIES:
                category = 'Other'
            confidence = str(d.get('confidence', 'low')).lower()
            if confidence not in ('high', 'medium', 'low'):
                confidence = 'low'
            impact = str(d.get('impact', 'needs_verification')).lower()
            if impact not in ('action_required', 'no_significant_impact', 'needs_verification'):
                impact = 'needs_verification'
            priority = str(d.get('priority', 'low')).lower()
            if priority not in ('critical', 'high', 'medium', 'low'):
                priority = 'low'
            visibility = str(d.get('visibility', 'limited')).lower()
            if visibility not in ('clear', 'limited'):
                visibility = 'limited'
            action = str(d.get('action', '')).strip()
            # Conservative override for far-distance photos:
            # do not present low-confidence or limited-visibility findings as confirmed actions.
            if distance == 'far' and (visibility != 'clear' or confidence == 'low'):
                impact = 'needs_verification'
                action = 'Obtain close-up photo / site verification'
                priority = 'low'
            if impact == 'no_significant_impact':
                action = 'No immediate action; routine monitoring'
                priority = 'low'
            if impact == 'needs_verification' and not action:
                action = 'Obtain close-up photo / site verification'
            out.append({
                'Category': category, 'Defect': defect, 'Visual Evidence': evidence,
                'Impact': impact, 'Priority': priority, 'Corrective Action': action,
                'Confidence': confidence, 'Visibility': visibility,
                'Photo Distance': distance,
                'Review Status': 'Pending engineer review',
            })
        return out, context, ''
    except Exception as exc:
        return [], {}, f'Invalid JSON: {str(exc)[:160]}'


def cached_result(key):
    if CLOUD_DEPLOYMENT:
        return None
    p = CACHE / f'{key}.json'
    try:
        return json.loads(p.read_text(encoding='utf-8')) if p.exists() else None
    except Exception:
        return None


def save_result(key, result):
    if CLOUD_DEPLOYMENT:
        return
    p, tmp = CACHE / f'{key}.json', CACHE / f'{key}.tmp'
    tmp.write_text(json.dumps(result, ensure_ascii=False), encoding='utf-8')
    tmp.replace(p)


def report_bytes(rows, raw_rows, texts, metadata):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Defects'
    active = [r for r in rows if r.get('Review Status') != 'Rejected']
    action_rows = [r for r in active if r.get('Impact') == 'action_required']
    verify_rows = [r for r in active if r.get('Impact') == 'needs_verification']
    minor_rows = [r for r in active if r.get('Impact') == 'no_significant_impact']
    sheets = [(ws, rows), (wb.create_sheet('Action Required'), action_rows),
              (wb.create_sheet('Needs Verification'), verify_rows),
              (wb.create_sheet('No Significant Impact'), minor_rows),
              (wb.create_sheet('Raw AI Responses'), raw_rows),
              (wb.create_sheet('PDF Text'), texts), (wb.create_sheet('Metadata'), metadata)]
    for sheet, records in sheets:
        if not records:
            sheet.append(['No data'])
            continue
        headers = list(records[0].keys())
        sheet.append(headers)
        for record in records:
            sheet.append([str(record.get(h, ''))[:32000] for h in headers])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='153A58')
            cell.font = Font(color='FFFFFF', bold=True)
        for j, h in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(j)].width = min(58, max(15, len(h) + 3))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


st.set_page_config(page_title='Telecom AI Inspector V4', page_icon='📡', layout='wide')
st.title('📡 Telecom AI Inspector V4')
st.caption('Action-first review • Far-photo caution • RunPod Cloud • Excel')

# Require an app password before rendering uploads, findings or API settings.
import hmac
try:
    APP_PASSWORD = st.secrets["APP_PASSWORD"]
    RUNPOD_BASE_URL = st.secrets["RUNPOD_BASE_URL"]
    RUNPOD_VLLM_API_KEY = st.secrets["RUNPOD_VLLM_API_KEY"]
except (KeyError, FileNotFoundError):
    st.error("Missing cloud secrets: APP_PASSWORD, RUNPOD_BASE_URL, RUNPOD_VLLM_API_KEY.")
    st.stop()

if not st.session_state.get("authenticated", False):
    with st.form("login_form"):
        password = st.text_input("Application password", type="password")
        submitted = st.form_submit_button("Sign in")
    if submitted:
        if hmac.compare_digest(password, str(APP_PASSWORD)):
            st.session_state["authenticated"] = True
            st.rerun()
        else:
            st.error("Incorrect password")
    st.stop()

with st.sidebar:
    if st.button("Sign out"):
        st.session_state.clear()
        st.rerun()
    st.header('Configuration')
    site = st.text_input('Site ID', 'BAG0906').strip().upper()
    provider = 'RunPod Cloud'
    model = CLOUD_MODEL
    url = str(RUNPOD_BASE_URL).strip()
    api_key = str(RUNPOD_VLLM_API_KEY)
    st.caption('Provider: RunPod Cloud (configured securely in Streamlit secrets).')
    timeout = st.number_input('Timeout per image (seconds)', min_value=60, max_value=1800, value=600, step=60)
    max_side = st.select_slider('Max image dimension', options=[640, 960, 1280, 1600], value=960)
    max_images = st.number_input('Maximum images per PDF', min_value=1, max_value=300, value=30)
    min_side = st.number_input('Minimum image side (pixels)', min_value=50, max_value=800, value=220)
    min_kb = st.number_input('Minimum embedded image size (KB)', min_value=1, max_value=500, value=15)
    render_pages = st.checkbox('Analyze rendered PDF pages instead of embedded photos (for scans)', value=False)
    st.caption('Cloud mode: analysis cache is disabled; results are held in the active session only.')

uploads = st.file_uploader('Upload JPG, PNG, WEBP or PDF',
                           type=['jpg', 'jpeg', 'png', 'webp', 'pdf'], accept_multiple_files=True)
st.warning('AI output is preliminary. Review defects before maintenance decisions. Confirm Site ID.')
cloud_approved = True
if provider == 'RunPod Cloud':
    cloud_approved = st.checkbox('I am authorized to send these telecom site images to the RunPod cloud service.', value=False)

if uploads:
    pdf_names = ' '.join(f.name for f in uploads if f.name.lower().endswith('.pdf'))
    found = sorted(set(m.upper() for m in SITE_RE.findall(pdf_names)))
    if found and site not in found:
        st.error(f'Site ID mismatch: selected {site}, filename suggests {", ".join(found)}.')
    elif found:
        st.success(f'Site ID matches filename: {site}')

if st.button('Analyze / Resume', type='primary', disabled=not uploads or not cloud_approved):
    if not re.fullmatch(r'[A-Z0-9_-]{3,32}', site):
        st.error('Invalid Site ID')
        st.stop()
    if not model or not url:
        st.error('Model and server URL are required.')
        st.stop()
    if provider == 'RunPod Cloud' and (not api_key or not url.startswith('https://')):
        st.error('RunPod requires an HTTPS base URL and the configured vLLM API key.')
        st.stop()
    pdf_names = ' '.join(f.name for f in uploads if f.name.lower().endswith('.pdf'))
    found = sorted(set(m.upper() for m in SITE_RE.findall(pdf_names)))
    if found and site not in found:
        st.error('Please correct Site ID first.')
        st.stop()
    jobs, texts, notes = [], [], []
    for f in uploads:
        data = f.getvalue()
        if len(data) > 100 * 1024 * 1024:
            notes.append(f'{f.name}: skipped (over 100MB)')
            continue
        try:
            if f.name.lower().endswith('.pdf'):
                jj, tt, nn = pdf_jobs(data, f.name, int(max_images), int(min_side),
                                      int(min_kb), render_pages, int(max_side))
                jobs += jj
                texts += tt
                notes += nn
            else:
                jobs.append((f.name, prepared(data, int(max_side)), None))
        except Exception as exc:
            notes.append(f'{f.name}: {exc}')
    for note in notes:
        st.info(note)
    if not jobs:
        st.warning('No photos found. For scanned PDF, enable rendered PDF pages.')
        st.stop()
    results = []
    bar = st.progress(0)
    summary = st.empty()
    start_all = time.perf_counter()
    for i, (source, img, page) in enumerate(jobs):
        # Include provider and URL, so local/cloud results cannot collide.
        cache_identity = json.dumps({'provider': provider, 'model': model, 'url': url.rstrip('/'),
                                     'prompt': PROMPT, 'version': 'v4', 'max_side': max_side}, sort_keys=True)
        key = hashlib.sha256(cache_identity.encode() + img).hexdigest()
        result = cached_result(key)
        was_cached = result is not None
        if result is None:
            started = time.perf_counter()
            try:
                if provider == 'Local Ollama':
                    raw_response, engine_seconds = call_ollama(img, model, url, int(timeout))
                else:
                    raw_response, engine_seconds = call_runpod(img, model, url, api_key, int(timeout))
                defects, context, warning = parse(raw_response)
                result = {'raw': raw_response, 'defects': defects, 'context': context, 'warning': warning,
                          'seconds': round(time.perf_counter() - started, 2),
                          'engine_seconds': round(engine_seconds, 2)}
                # Invalid responses are not cached; they can be retried after fixing the model.
                if not warning:
                    save_result(key, result)
            except Exception as exc:
                msg = str(exc)
                # Do not expose an accidentally embedded API key in error messages.
                if api_key:
                    msg = msg.replace(api_key, '[REDACTED]')
                result = {'raw': '', 'defects': [], 'warning': f'API error: {msg[:400]}',
                          'seconds': round(time.perf_counter() - started, 2)}
        results.append({'source': source, 'page': page, 'result': result, 'cached': was_cached})
        bar.progress((i + 1) / len(jobs))
        summary.write(f'Completed {i+1}/{len(jobs)} images; elapsed {round((time.perf_counter()-start_all)/60, 1)} min')
    st.session_state['v4_results'] = results
    st.session_state['v4_texts'] = texts
    st.session_state['v4_site'] = site
    st.session_state['v4_provider'] = provider
    st.session_state['v4_model'] = model

if 'v4_results' in st.session_state:
    results = st.session_state['v4_results']
    report_site = st.session_state['v4_site']
    defects, raw_rows = [], []
    for entry in results:
        result = entry['result']
        raw_rows.append({'Source': entry['source'], 'Page': entry['page'] or '',
                         'Raw Response': result.get('raw', ''),
                         'Warning': result.get('warning', ''),
                         'Seconds': result.get('seconds', ''), 'Cache Hit': entry['cached'],
                         'Photo Distance': result.get('context', {}).get('Photo Distance', 'unknown'),
                         'Distance Reason': result.get('context', {}).get('Distance Reason', ''),
                         'Uncertain Observations': result.get('context', {}).get('Uncertain Observations', '')})
        for d in result.get('defects', []):
            defects.append({'Site ID': report_site, 'Source': entry['source'],
                            'Page': entry['page'] or '', **d,
                            'Seconds': result.get('seconds', '')})
    st.subheader('Action-focused inspection')
    errors = sum(bool(r['Warning']) for r in raw_rows)
    far_count = sum(r['Photo Distance'] == 'far' for r in raw_rows)
    action_count = sum(d['Impact'] == 'action_required' for d in defects)
    minor_count = sum(d['Impact'] == 'no_significant_impact' for d in defects)
    verify_count = sum(d['Impact'] == 'needs_verification' for d in defects)
    st.write(f"Images: {len(results)} | Action required: {action_count} | No significant impact: {minor_count} | Verify: {verify_count} | Far photos: {far_count} | AI errors: {errors}")
    if errors:
        st.error(f'{errors} image(s) failed or returned invalid JSON. No defects does not mean a passed inspection.')
    if far_count:
        st.warning(f'{far_count} photo(s) appear to be taken from far away. Small equipment details cannot be verified; request close-ups where needed.')
    if defects:
        st.caption('AI impact and priority are preliminary. Edit them after reviewing the original images.')
        reviewed = st.data_editor(
            pd.DataFrame(defects), num_rows='dynamic', hide_index=True, use_container_width=True,
            column_config={
                'Review Status': st.column_config.SelectboxColumn('Review Status',
                    options=['Pending engineer review', 'Confirmed', 'Rejected', 'Needs site verification']),
                'Impact': st.column_config.SelectboxColumn('Impact',
                    options=['action_required', 'no_significant_impact', 'needs_verification']),
                'Priority': st.column_config.SelectboxColumn('Priority',
                    options=['critical', 'high', 'medium', 'low']),
            }, key='review_editor_v4')
        approved = reviewed.to_dict('records')
        active = [d for d in approved if d.get('Review Status') != 'Rejected']
        action_items = [d for d in active if d.get('Impact') == 'action_required']
        verify_items = [d for d in active if d.get('Impact') == 'needs_verification']
        minor_items = [d for d in active if d.get('Impact') == 'no_significant_impact']
        st.markdown('#### Action required')
        if action_items:
            st.dataframe(pd.DataFrame(action_items), hide_index=True, use_container_width=True)
        else:
            st.info('No action-required findings proposed.')
        st.markdown('#### Needs close-up / site verification')
        if verify_items:
            st.dataframe(pd.DataFrame(verify_items), hide_index=True, use_container_width=True)
        else:
            st.caption('No verification-only findings proposed.')
        with st.expander(f'No significant impact — {len(minor_items)} minor observations'):
            if minor_items:
                st.dataframe(pd.DataFrame(minor_items), hide_index=True, use_container_width=True)
            else:
                st.caption('None')
    else:
        st.info('No structured findings. Review raw AI responses and far-photo flags.')
        approved = []
    meta = [{'Key': 'Site ID', 'Value': report_site},
            {'Key': 'Provider', 'Value': st.session_state['v4_provider']},
            {'Key': 'Model', 'Value': st.session_state['v4_model']},
            {'Key': 'Generated', 'Value': datetime.now().isoformat(timespec='seconds')},
            {'Key': 'AI errors', 'Value': errors},
            {'Key': 'Notice', 'Value': 'All AI findings require engineer verification'}]
    meta += [{'Key': 'Action required (AI preliminary)', 'Value': action_count},
             {'Key': 'No significant impact (AI preliminary)', 'Value': minor_count},
             {'Key': 'Needs verification (AI preliminary)', 'Value': verify_count},
             {'Key': 'Far photos (AI estimated)', 'Value': far_count}]
    file_bytes = report_bytes(approved, raw_rows, st.session_state['v4_texts'], meta)
    st.download_button('⬇️ Download reviewed Excel',
                       file_name=f'{report_site}_inspection_V4_{datetime.now():%Y%m%d_%H%M}.xlsx',
                       data=file_bytes,
                       mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    with st.expander('Raw model responses and warnings'):
        st.dataframe(pd.DataFrame(raw_rows), use_container_width=True)
